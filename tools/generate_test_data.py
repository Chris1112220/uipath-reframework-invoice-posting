"""Generate the fictional sample data used by the Invoice Posting Bot.

Every company, vendor, PO, GL account and amount in these files is made up.
Re-run any time to reset Data/ and Tests/TestData/ to a known state:

    pip install openpyxl
    python tools/generate_test_data.py
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = ROOT / "Data" / "Input"
TEST_DATA_DIR = ROOT / "Tests" / "TestData"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")

# ---------------------------------------------------------------------------
# Reference data (fictional)
# ---------------------------------------------------------------------------
VENDORS = [
    # VendorId, VendorName, Status, DefaultExpenseGL, GLDescription, PaymentTerms
    ("V001", "Northwind Supply Co.", "Active", "6100-00", "Office Supplies", "Net 30"),
    ("V002", "Contoso Office Furniture", "Active", "6200-00", "Furniture & Equipment", "Net 45"),
    ("V003", "Fabrikam Facilities Services", "Active", "6300-00", "Facilities Maintenance", "Net 30"),
    ("V004", "Tailspin Logistics", "Active", "6400-00", "Freight & Delivery", "Net 15"),
    ("V005", "Adventure Works Printing", "Inactive", "6500-00", "Printing & Reproduction", "Net 30"),
    ("V006", "Litware IT Services", "Active", "6600-00", "IT Services", "Net 30"),
]

PURCHASE_ORDERS = [
    # PONumber, VendorId, POAmount, ReceivedAmount, Status
    ("PO-5001", "V001", 1250.00, 1250.00, "Open"),
    ("PO-5002", "V002", 4800.00, 4800.00, "Open"),
    ("PO-5003", "V003", 2150.75, 2150.75, "Open"),
    ("PO-5004", "V004", 980.40, 980.40, "Open"),
    ("PO-5005", "V001", 3000.00, 3000.00, "Open"),
    ("PO-5006", "V004", 1000.00, 1000.00, "Open"),
    ("PO-5007", "V006", 31500.00, 31500.00, "Open"),
    ("PO-5008", "V006", 600.00, 300.00, "Open"),  # only half the goods received
    ("PO-5009", "V003", 2000.00, 2000.00, "Open"),
    ("PO-5010", "V005", 750.00, 750.00, "Open"),
]

# InvoiceNumber, InvoiceDate, VendorId, VendorName, PONumber, Amount, Currency, Description, Scenario
INVOICES = [
    ("INV-10001", "2026-09-01", "V001", "Northwind Supply Co.", "PO-5001", 1250.00, "USD", "Copy paper and toner", "Happy path"),
    ("INV-10002", "2026-09-02", "V002", "Contoso Office Furniture", "PO-5002", 4800.00, "USD", "Ergonomic chairs (12)", "Happy path"),
    ("INV-10003", "2026-09-02", "V003", "Fabrikam Facilities Services", "PO-5003", 2150.75, "USD", "HVAC quarterly service", "Happy path"),
    ("INV-10004", "2026-09-03", "V004", "Tailspin Logistics", "PO-5004", 980.40, "USD", "Inbound freight", "Happy path"),
    ("INV-10005", "2026-09-03", "V001", "Northwind Supply Co.", "PO-5005", 3020.00, "USD", "Breakroom supplies", "Within 2% tolerance (+0.67%)"),
    ("INV-10002", "2026-09-04", "V002", "Contoso Office Furniture", "PO-5002", 4800.00, "USD", "Ergonomic chairs (12)", "Duplicate invoice number"),
    ("INV-10006", "2026-09-04", "V002", "Contoso Office Furniture", "", 615.00, "USD", "Desk lamps", "Missing PO"),
    ("INV-10007", "2026-09-05", "V999", "Unknown Vendor LLC", "PO-5001", 300.00, "USD", "Consulting", "Unknown vendor"),
    ("INV-10008", "2026-09-05", "V003", "Fabrikam Facilities Services", "PO-5003", -450.00, "USD", "Credit memo sent as invoice", "Negative amount"),
    ("INV-10009", "2026-09-08", "V004", "Tailspin Logistics", "PO-5006", 1150.00, "USD", "Expedited freight", "Over tolerance (+15%)"),
    ("INV-10010", "2026-09-08", "V005", "Adventure Works Printing", "PO-5010", 750.00, "USD", "Brochure printing", "Inactive vendor"),
    ("INV-10011", "2026-09-09", "V001", "Northwind Supply Co.", "PO-5002", 4800.00, "USD", "Office supplies", "PO belongs to another vendor"),
    ("INV-10012", "2026-09-09", "V006", "Litware IT Services", "PO-5007", 31500.00, "USD", "Annual software support", "Above auto-post limit"),
    ("INV-10013", "2026-09-10", "V006", "Litware IT Services", "PO-5008", 600.00, "USD", "Laptop docks", "3-way match: not fully received"),
    ("INV-10014", "2026-09-10", "V003", "Fabrikam Facilities Services", "PO-5009", 1999.99, "USD", "Window cleaning", "Happy path"),
    ("", "2026-09-11", "V001", "Northwind Supply Co.", "PO-5001", 99.00, "USD", "Pens", "Missing invoice number"),
    ("INV-10015", "2026-09-11", "V004", "Tailspin Logistics", "PO-5004", "TBD", "USD", "Freight (amount pending)", "Non-numeric amount"),
]

CONFIG_SETTINGS = [
    ("logF_BusinessProcessName", "Finance_InvoicePosting", "Business process name added to every log line (REFramework convention)."),
    ("RunMode", "Local", "Local = read Data/Input/Invoices.xlsx directly (no Orchestrator needed). Orchestrator = use the InvoicePosting queue."),
    ("OrchestratorQueueName", "InvoicePosting", "Queue the Dispatcher fills and the Performer consumes. Enable 'Enforce unique references' on it."),
    ("OrchestratorQueueFolder", "", "Orchestrator folder for the queue and assets. Blank = the robot's default folder."),
    ("InvoiceInputPath", r"Data\Input\Invoices.xlsx", "Vendor invoice register read by the Dispatcher."),
    ("InvoiceSheetName", "Invoices", "Sheet name in the invoice register."),
    ("VendorMasterPath", r"Data\Input\VendorMaster.xlsx", "Vendor master (vendor status and default expense GL)."),
    ("VendorSheetName", "Vendors", "Sheet name in the vendor master."),
    ("PurchaseOrdersPath", r"Data\Input\PurchaseOrders.xlsx", "Open purchase orders with received amounts (for 3-way match)."),
    ("PurchaseOrderSheetName", "PurchaseOrders", "Sheet name in the PO file."),
    ("JournalOutputPath", r"Data\Output\JournalEntries.xlsx", "Simulated ERP: journal lines are appended here."),
    ("JournalSheetName", "JournalEntries", "Sheet name in the journal output."),
    ("RunReportFolder", r"Data\Output", "Folder for the per-run RunReport_*.xlsx and DispatcherReport_*.xlsx."),
    ("ProcessesToKill", "", "Comma-separated process names KillAllProcesses closes (blank = none; this bot only touches files)."),
]

CONFIG_CONSTANTS = [
    ("MaxRetryNumber", 2, "Local mode retries for system exceptions. Keep 0 in Orchestrator mode and use the queue's Auto Retry instead."),
    ("MaxConsecutiveSystemExceptions", 3, "Stop the run after this many system exceptions in a row (0 = disabled)."),
    ("APAccount", "2000-00", "Accounts Payable control account credited for every invoice."),
    ("JournalPrefix", "AP", "Prefix for generated journal IDs (AP-<InvoiceNumber>)."),
    ("ExScreenshotsFolderPath", "Exceptions_Screenshots", "Where screenshots are saved when a system exception occurs."),
    ("LogMessage_GetTransactionData", "Processing Transaction Number: ", "Standard REFramework log prefix."),
    ("LogMessage_GetTransactionDataError", "Error getting transaction data for Transaction Number: ", "Standard REFramework log prefix."),
    ("LogMessage_Success", "Transaction Successful.", "Standard REFramework log message."),
    ("LogMessage_BusinessRuleException", "Business rule exception: ", "Standard REFramework log prefix."),
    ("LogMessage_ApplicationException", "System exception: ", "Standard REFramework log prefix."),
    ("ExceptionMessage_ConsecutiveErrors", "The maximum number of consecutive system exceptions was reached.", "Message used to stop the run."),
]

CONFIG_ASSETS = [
    # Name, Asset, OrchestratorAssetFolder, LocalValue, Description
    ("MatchTolerancePercent", "InvoicePosting_MatchTolerancePercent", "", 2,
     "Allowed % difference between invoice and PO / received amount. LocalValue is used when the asset can't be read."),
    ("MaxAutoPostAmount", "InvoicePosting_MaxAutoPostAmount", "", 25000,
     "Invoices above this amount need manual approval and are not auto-posted."),
]

# Test cases for Tests/Test_ValidateInvoice_DataDriven.xaml
VALIDATION_CASES = [
    # CaseId, VendorId, PONumber, Amount, ExpectedResult, ExpectedMessageContains
    ("TC01", "V001", "PO-5001", 1250.00, "Pass", ""),
    ("TC02", "V001", "PO-5005", 3020.00, "Pass", ""),
    ("TC03", "V001", "PO-5005", 3060.00, "Pass", ""),
    ("TC04", "V001", "PO-5005", 3061.00, "BusinessRuleException", "tolerance"),
    ("TC05", "V999", "PO-5001", 100.00, "BusinessRuleException", "Unknown vendor"),
    ("TC06", "V005", "PO-5010", 750.00, "BusinessRuleException", "inactive"),
    ("TC07", "V002", "", 615.00, "BusinessRuleException", "Missing PO"),
    ("TC08", "V002", "PO-9999", 615.00, "BusinessRuleException", "not found"),
    ("TC09", "V001", "PO-5002", 4800.00, "BusinessRuleException", "belongs to vendor"),
    ("TC10", "V003", "PO-5003", 0.00, "BusinessRuleException", "greater than zero"),
    ("TC11", "V003", "PO-5003", -450.00, "BusinessRuleException", "greater than zero"),
    ("TC12", "V006", "PO-5008", 600.00, "BusinessRuleException", "goods received"),
    ("TC13", "V006", "PO-5007", 31500.00, "BusinessRuleException", "auto-post limit"),
]


def write_sheet(wb: Workbook, title: str, headers: list[str], rows: list[tuple], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = "A2"
    for idx, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[idx - 1])) for r in rows]) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(width, 70)
    return ws


def money_format(ws, columns: list[str]):
    headers = [c.value for c in ws[1]]
    for name in columns:
        col = headers.index(name) + 1
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"


def main() -> None:
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    TEST_DATA_DIR.mkdir(parents=True, exist_ok=True)

    # Invoices.xlsx (Scenario column is documentation only; the bot ignores it)
    wb = Workbook()
    ws = write_sheet(wb, "Invoices",
                     ["InvoiceNumber", "InvoiceDate", "VendorId", "VendorName", "PONumber",
                      "Amount", "Currency", "Description", "Scenario"], INVOICES, first=True)
    money_format(ws, ["Amount"])
    wb.save(INPUT_DIR / "Invoices.xlsx")

    wb = Workbook()
    write_sheet(wb, "Vendors",
                ["VendorId", "VendorName", "Status", "DefaultExpenseGL", "GLDescription", "PaymentTerms"],
                VENDORS, first=True)
    wb.save(INPUT_DIR / "VendorMaster.xlsx")

    wb = Workbook()
    ws = write_sheet(wb, "PurchaseOrders",
                     ["PONumber", "VendorId", "POAmount", "ReceivedAmount", "Status"],
                     PURCHASE_ORDERS, first=True)
    money_format(ws, ["POAmount", "ReceivedAmount"])
    wb.save(INPUT_DIR / "PurchaseOrders.xlsx")

    wb = Workbook()
    write_sheet(wb, "Settings", ["Name", "Value", "Description"], CONFIG_SETTINGS, first=True)
    write_sheet(wb, "Constants", ["Name", "Value", "Description"], CONFIG_CONSTANTS)
    write_sheet(wb, "Assets", ["Name", "Asset", "OrchestratorAssetFolder", "LocalValue", "Description"],
                CONFIG_ASSETS)
    wb.save(ROOT / "Data" / "Config.xlsx")

    wb = Workbook()
    ws = write_sheet(wb, "Cases",
                     ["CaseId", "VendorId", "PONumber", "Amount", "ExpectedResult", "ExpectedMessageContains"],
                     VALIDATION_CASES, first=True)
    money_format(ws, ["Amount"])
    wb.save(TEST_DATA_DIR / "ValidateInvoice_Cases.xlsx")

    print(f"Wrote {len(INVOICES)} invoices, {len(VENDORS)} vendors, {len(PURCHASE_ORDERS)} POs, "
          f"{len(VALIDATION_CASES)} validation test cases.")


if __name__ == "__main__":
    main()
